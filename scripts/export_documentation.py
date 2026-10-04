"""Export Docs/OFFICIAL_PROJECT_DOCUMENTATION.md to both DOCX (Word) and PDF formats.

Produces:
- Docs/OFFICIAL_PROJECT_DOCUMENTATION.docx
- Docs/OFFICIAL_PROJECT_DOCUMENTATION.pdf
"""

from __future__ import annotations

import os
import re
import subprocess
import sys
from pathlib import Path

# python-docx imports
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

ROOT_DIR = Path(__file__).resolve().parent.parent
DOC_MD_PATH = ROOT_DIR / "Docs" / "OFFICIAL_PROJECT_DOCUMENTATION.md"
DOCX_OUT_PATH = ROOT_DIR / "Docs" / "OFFICIAL_PROJECT_DOCUMENTATION.docx"
PDF_OUT_PATH = ROOT_DIR / "Docs" / "OFFICIAL_PROJECT_DOCUMENTATION.pdf"
HTML_TEMP_PATH = ROOT_DIR / "Docs" / "_temp_doc_for_pdf.html"


def set_cell_background(cell, fill_hex: str):
    """Set the background color of a table cell."""
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)


def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    """Set inner margins (padding) of a table cell in dxa (1 pt = 20 dxa)."""
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(
        f'<w:tcMar {nsdecls("w")}>'
        f'<w:top w:w="{top}" w:type="dxa"/>'
        f'<w:bottom w:w="{bottom}" w:type="dxa"/>'
        f'<w:left w:w="{left}" w:type="dxa"/>'
        f'<w:right w:w="{right}" w:type="dxa"/>'
        f'</w:tcMar>'
    )
    tcPr.append(tcMar)


def set_table_borders(table, color="D0D7DE"):
    """Set thin clean borders on a table."""
    tblPr = table._tbl.tblPr
    borders = parse_xml(
        f'<w:tblBorders {nsdecls("w")}>'
        f'<w:top w:val="single" w:sz="4" w:space="0" w:color="{color}"/>'
        f'<w:bottom w:val="single" w:sz="6" w:space="0" w:color="{color}"/>'
        f'<w:insideH w:val="single" w:sz="4" w:space="0" w:color="{color}"/>'
        f'<w:insideV w:val="none"/>'
        f'<w:left w:val="none"/>'
        f'<w:right w:val="none"/>'
        f'</w:tblBorders>'
    )
    tblPr.append(borders)


def add_formatted_runs(paragraph, text: str, default_font_size=Pt(10.5), is_italic=False):
    """Parse inline markdown (bold, italic, inline code) and add runs to paragraph."""
    # Pattern to match bold `**text**`, inline code `` `text` ``, and italic `*text*`
    pattern = re.compile(r'(\*\*.*?\*\*|`.*?`|\*.*?\*)')
    parts = pattern.split(text)
    
    for part in parts:
        if not part:
            continue
        if part.startswith('**') and part.endswith('**') and len(part) >= 4:
            run = paragraph.add_run(part[2:-2])
            run.bold = True
            run.font.size = default_font_size
            if is_italic:
                run.italic = True
        elif part.startswith('`') and part.endswith('`') and len(part) >= 2:
            run = paragraph.add_run(part[1:-1])
            run.font.name = "Consolas"
            run.font.size = Pt(9.5)
            run.font.color.rgb = RGBColor(176, 32, 55)  # Markdown red/purple accent
        elif part.startswith('*') and part.endswith('*') and len(part) >= 2:
            run = paragraph.add_run(part[1:-1])
            run.italic = True
            run.font.size = default_font_size
        else:
            run = paragraph.add_run(part)
            run.font.size = default_font_size
            if is_italic:
                run.italic = True


def build_word_document(md_text: str, output_path: Path):
    """Parse Markdown content and build a Word (.docx) document."""
    doc = docx.Document()

    # Configure Margins: 1 inch all sides
    for section in doc.sections:
        section.top_margin = Inches(1.0)
        section.bottom_margin = Inches(1.0)
        section.left_margin = Inches(1.0)
        section.right_margin = Inches(1.0)

    # Base Style Customization
    normal_style = doc.styles['Normal']
    normal_style.font.name = 'Segoe UI'
    normal_style.font.size = Pt(10.5)
    normal_style.font.color.rgb = RGBColor(36, 41, 47)  # #24292f
    normal_style.paragraph_format.line_spacing = 1.25
    normal_style.paragraph_format.space_after = Pt(6)

    lines = md_text.splitlines()
    i = 0
    n = len(lines)

    while i < n:
        line = lines[i]
        stripped = line.strip()

        # Blank line
        if not stripped:
            i += 1
            continue

        # Horizontal rule
        if stripped in ("---", "***", "___"):
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(8)
            p.paragraph_format.space_after = Pt(8)
            pBdr = parse_xml(f'<w:pBdr {nsdecls("w")}><w:bottom w:val="single" w:sz="6" w:space="1" w:color="E1E4E8"/></w:pBdr>')
            p._p.get_or_add_pPr().append(pBdr)
            i += 1
            continue

        # Code block
        if stripped.startswith("```"):
            code_lines = []
            i += 1
            while i < n and not lines[i].strip().startswith("```"):
                code_lines.append(lines[i])
                i += 1
            i += 1  # consume closing ```
            
            code_text = "\n".join(code_lines)
            table = doc.add_table(rows=1, cols=1)
            table.alignment = WD_TABLE_ALIGNMENT.CENTER
            cell = table.cell(0, 0)
            set_cell_background(cell, "F6F8FA")
            set_cell_margins(cell, top=140, bottom=140, left=200, right=200)
            
            # Left border styling for code box
            tcPr = cell._tc.get_or_add_tcPr()
            borders = parse_xml(
                f'<w:tcBorders {nsdecls("w")}>'
                f'<w:left w:val="single" w:sz="16" w:space="0" w:color="0366D6"/>'
                f'<w:top w:val="single" w:sz="4" w:space="0" w:color="E1E4E8"/>'
                f'<w:right w:val="single" w:sz="4" w:space="0" w:color="E1E4E8"/>'
                f'<w:bottom w:val="single" w:sz="4" w:space="0" w:color="E1E4E8"/>'
                f'</w:tcBorders>'
            )
            tcPr.append(borders)
            
            cp = cell.paragraphs[0]
            cp.paragraph_format.space_after = Pt(0)
            cp.paragraph_format.line_spacing = 1.15
            run = cp.add_run(code_text)
            run.font.name = "Consolas"
            run.font.size = Pt(8.5)
            run.font.color.rgb = RGBColor(36, 41, 47)
            
            doc.add_paragraph().paragraph_format.space_after = Pt(4)
            continue

        # Headings
        if stripped.startswith("# ") and not stripped.startswith("## "):
            text = stripped[2:].strip()
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(16)
            p.paragraph_format.space_after = Pt(6)
            run = p.add_run(text)
            run.bold = True
            run.font.size = Pt(22)
            run.font.name = "Segoe UI Semibold"
            run.font.color.rgb = RGBColor(9, 105, 218)  # #0969da
            i += 1
            continue

        if stripped.startswith("## "):
            text = stripped[3:].strip()
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(14)
            p.paragraph_format.space_after = Pt(4)
            run = p.add_run(text)
            run.bold = True
            run.font.size = Pt(15)
            run.font.name = "Segoe UI Semibold"
            run.font.color.rgb = RGBColor(31, 35, 40)
            i += 1
            continue

        if stripped.startswith("### "):
            text = stripped[4:].strip()
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(10)
            p.paragraph_format.space_after = Pt(2)
            run = p.add_run(text)
            run.bold = True
            run.font.size = Pt(12)
            run.font.name = "Segoe UI Semibold"
            run.font.color.rgb = RGBColor(87, 96, 106)
            i += 1
            continue

        # Blockquote
        if stripped.startswith(">"):
            quote_text = stripped.lstrip("> ").strip()
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Inches(0.3)
            p.paragraph_format.space_before = Pt(4)
            p.paragraph_format.space_after = Pt(6)
            pPr = p._p.get_or_add_pPr()
            pBdr = parse_xml(f'<w:pBdr {nsdecls("w")}><w:left w:val="single" w:sz="18" w:space="8" w:color="0969DA"/></w:pBdr>')
            pPr.append(pBdr)
            add_formatted_runs(p, quote_text, default_font_size=Pt(10), is_italic=True)
            i += 1
            continue

        # Markdown Table
        if stripped.startswith("|") and stripped.endswith("|") and i + 1 < n and "|---" in lines[i + 1]:
            table_lines = [stripped]
            i += 1
            separator_line = lines[i].strip()
            i += 1
            while i < n and lines[i].strip().startswith("|") and lines[i].strip().endswith("|"):
                table_lines.append(lines[i].strip())
                i += 1

            headers = [c.strip() for c in table_lines[0].split("|")[1:-1]]
            num_cols = len(headers)
            
            table = doc.add_table(rows=len(table_lines), cols=num_cols)
            table.alignment = WD_TABLE_ALIGNMENT.CENTER
            set_table_borders(table)

            # Style Header Row
            for col_idx, h_text in enumerate(headers):
                cell = table.cell(0, col_idx)
                set_cell_background(cell, "F2F4F8")
                set_cell_margins(cell, top=100, bottom=100, left=140, right=140)
                cp = cell.paragraphs[0]
                cp.paragraph_format.space_after = Pt(2)
                add_formatted_runs(cp, f"**{h_text}**", default_font_size=Pt(9.5))

            # Style Data Rows
            for row_idx, r_line in enumerate(table_lines[1:], start=1):
                cols = [c.strip() for c in r_line.split("|")[1:-1]]
                for col_idx in range(min(num_cols, len(cols))):
                    cell = table.cell(row_idx, col_idx)
                    if row_idx % 2 == 0:
                        set_cell_background(cell, "FAFBFC")
                    set_cell_margins(cell, top=80, bottom=80, left=140, right=140)
                    cp = cell.paragraphs[0]
                    cp.paragraph_format.space_after = Pt(2)
                    add_formatted_runs(cp, cols[col_idx], default_font_size=Pt(9))

            doc.add_paragraph().paragraph_format.space_after = Pt(4)
            continue

        # Unordered List (* or -)
        if stripped.startswith(("* ", "- ")):
            item_text = stripped[2:].strip()
            p = doc.add_paragraph(style='List Bullet')
            p.paragraph_format.space_after = Pt(3)
            p.paragraph_format.space_before = Pt(1)
            add_formatted_runs(p, item_text, default_font_size=Pt(10))
            i += 1
            continue

        # Numbered List (1. , 2. )
        num_match = re.match(r'^(\d+)\.\s+(.*)$', stripped)
        if num_match:
            item_text = num_match.group(2)
            p = doc.add_paragraph(style='List Number')
            p.paragraph_format.space_after = Pt(3)
            p.paragraph_format.space_before = Pt(1)
            add_formatted_runs(p, item_text, default_font_size=Pt(10))
            i += 1
            continue

        # Regular Paragraph
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(2)
        p.paragraph_format.space_after = Pt(6)
        add_formatted_runs(p, stripped, default_font_size=Pt(10))
        i += 1

    doc.save(str(output_path))
    print(f"Generated DOCX: {output_path} ({os.path.getsize(output_path):,} bytes)")


def build_html_for_pdf(md_text: str) -> str:
    """Transform markdown to publication-styled HTML ready for PDF rendering."""
    # Convert code blocks
    def code_replacer(m):
        lang = m.group(1) or ""
        code = m.group(2)
        escaped = (code.replace("&", "&amp;")
                       .replace("<", "&lt;")
                       .replace(">", "&gt;"))
        return f'<pre><code class="{lang}">{escaped}</code></pre>'
    
    # Process blocks
    lines = md_text.splitlines()
    html_parts = []
    i = 0
    n = len(lines)

    in_list = False
    list_type = None

    def close_list():
        nonlocal in_list, list_type
        if in_list:
            html_parts.append(f"</{list_type}>")
            in_list = False
            list_type = None

    while i < n:
        line = lines[i]
        s = line.strip()

        if not s:
            close_list()
            i += 1
            continue

        if s in ("---", "***", "___"):
            close_list()
            html_parts.append("<hr/>")
            i += 1
            continue

        if s.startswith("```"):
            close_list()
            code_lines = []
            i += 1
            while i < n and not lines[i].strip().startswith("```"):
                code_lines.append(lines[i])
                i += 1
            i += 1
            c_text = "\n".join(code_lines)
            escaped = (c_text.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;"))
            html_parts.append(f'<div class="code-container"><pre><code>{escaped}</code></pre></div>')
            continue

        if s.startswith("# ") and not s.startswith("## "):
            close_list()
            html_parts.append(f"<h1>{s[2:].strip()}</h1>")
            i += 1
            continue

        if s.startswith("## "):
            close_list()
            html_parts.append(f"<h2>{s[3:].strip()}</h2>")
            i += 1
            continue

        if s.startswith("### "):
            close_list()
            html_parts.append(f"<h3>{s[4:].strip()}</h3>")
            i += 1
            continue

        if s.startswith(">"):
            close_list()
            q_text = s.lstrip("> ").strip()
            q_text = re.sub(r'\*\*(.*?)\*\*', r'<strong>\1</strong>', q_text)
            q_text = re.sub(r'`(.*?)`', r'<code>\1</code>', q_text)
            q_text = re.sub(r'\*(.*?)\*', r'<em>\1</em>', q_text)
            html_parts.append(f"<blockquote>{q_text}</blockquote>")
            i += 1
            continue

        # Markdown Table
        if s.startswith("|") and s.endswith("|") and i + 1 < n and "|---" in lines[i + 1]:
            close_list()
            table_lines = [s]
            i += 2  # skip separator
            while i < n and lines[i].strip().startswith("|") and lines[i].strip().endswith("|"):
                table_lines.append(lines[i].strip())
                i += 1

            headers = [c.strip() for c in table_lines[0].split("|")[1:-1]]
            th_cells = "".join(f"<th>{h}</th>" for h in headers)
            rows = []
            for r_line in table_lines[1:]:
                cells = [c.strip() for c in r_line.split("|")[1:-1]]
                td_cells = []
                for cell in cells:
                    cell_html = re.sub(r'\*\*(.*?)\*\*', r'<strong>\1</strong>', cell)
                    cell_html = re.sub(r'`(.*?)`', r'<code>\1</code>', cell_html)
                    cell_html = re.sub(r'\*(.*?)\*', r'<em>\1</em>', cell_html)
                    td_cells.append(f"<td>{cell_html}</td>")
                rows.append(f"<tr>{''.join(td_cells)}</tr>")

            table_html = f"<table><thead><tr>{th_cells}</tr></thead><tbody>{''.join(rows)}</tbody></table>"
            html_parts.append(table_html)
            continue

        # Bullet list
        if s.startswith(("* ", "- ")):
            if not in_list or list_type != "ul":
                close_list()
                html_parts.append("<ul>")
                in_list = True
                list_type = "ul"
            item_text = s[2:].strip()
            item_html = re.sub(r'\*\*(.*?)\*\*', r'<strong>\1</strong>', item_text)
            item_html = re.sub(r'`(.*?)`', r'<code>\1</code>', item_html)
            item_html = re.sub(r'\*(.*?)\*', r'<em>\1</em>', item_html)
            html_parts.append(f"<li>{item_html}</li>")
            i += 1
            continue

        # Numbered list
        num_m = re.match(r'^(\d+)\.\s+(.*)$', s)
        if num_m:
            if not in_list or list_type != "ol":
                close_list()
                html_parts.append("<ol>")
                in_list = True
                list_type = "ol"
            item_text = num_m.group(2).strip()
            item_html = re.sub(r'\*\*(.*?)\*\*', r'<strong>\1</strong>', item_text)
            item_html = re.sub(r'`(.*?)`', r'<code>\1</code>', item_html)
            item_html = re.sub(r'\*(.*?)\*', r'<em>\1</em>', item_html)
            html_parts.append(f"<li>{item_html}</li>")
            i += 1
            continue

        # Regular Paragraph
        close_list()
        p_text = s
        p_html = re.sub(r'\*\*(.*?)\*\*', r'<strong>\1</strong>', p_text)
        p_html = re.sub(r'`(.*?)`', r'<code>\1</code>', p_html)
        p_html = re.sub(r'\*(.*?)\*', r'<em>\1</em>', p_html)
        html_parts.append(f"<p>{p_html}</p>")
        i += 1

    close_list()
    body_content = "\n".join(html_parts)

    html_template = f"""<!DOCTYPE html>
<html lang="en">
<head>
<meta charset="utf-8">
<title>Project Technical Specification & Official Documentation</title>
<style>
  @page {{
    size: A4;
    margin: 18mm 16mm 18mm 16mm;
    @bottom-right {{
      content: counter(page);
    }}
  }}
  body {{
    font-family: -apple-system, BlinkMacSystemFont, "Segoe UI", Roboto, "Helvetica Neue", Arial, sans-serif;
    color: #24292f;
    line-height: 1.5;
    font-size: 10pt;
    background: #ffffff;
    margin: 0;
    padding: 0;
  }}
  h1 {{
    font-size: 20pt;
    color: #0969da;
    border-bottom: 2px solid #0969da;
    padding-bottom: 8px;
    margin-top: 10px;
    margin-bottom: 12px;
    page-break-after: avoid;
  }}
  h2 {{
    font-size: 13.5pt;
    color: #1f2328;
    border-bottom: 1px solid #d0d7de;
    padding-bottom: 6px;
    margin-top: 22px;
    margin-bottom: 10px;
    page-break-after: avoid;
  }}
  h3 {{
    font-size: 11pt;
    color: #57606a;
    margin-top: 14px;
    margin-bottom: 6px;
    page-break-after: avoid;
  }}
  p {{
    margin-top: 4px;
    margin-bottom: 8px;
    text-align: justify;
  }}
  hr {{
    border: none;
    border-top: 1px solid #e1e4e8;
    margin: 16px 0;
  }}
  blockquote {{
    margin: 10px 0;
    padding: 8px 14px;
    background: #f6f8fa;
    border-left: 4px solid #0969da;
    color: #57606a;
    font-style: italic;
    page-break-inside: avoid;
  }}
  ul, ol {{
    margin: 4px 0 8px 18px;
    padding-left: 6px;
  }}
  li {{
    margin-bottom: 4px;
  }}
  code {{
    font-family: "Consolas", "Courier New", monospace;
    font-size: 8.5pt;
    background: #eff1f3;
    color: #b02037;
    padding: 2px 4px;
    border-radius: 4px;
  }}
  .code-container {{
    background: #f6f8fa;
    border: 1px solid #d0d7de;
    border-left: 4px solid #0969da;
    border-radius: 4px;
    padding: 10px 14px;
    margin: 10px 0;
    page-break-inside: avoid;
    overflow-x: auto;
  }}
  pre {{
    margin: 0;
    padding: 0;
  }}
  pre code {{
    background: transparent;
    padding: 0;
    color: #24292f;
    font-size: 8pt;
    line-height: 1.35;
    white-space: pre-wrap;
    word-break: break-all;
  }}
  table {{
    width: 100%;
    border-collapse: collapse;
    margin: 12px 0;
    font-size: 8.5pt;
    page-break-inside: avoid;
  }}
  th, td {{
    border: 1px solid #d0d7de;
    padding: 6px 10px;
    text-align: left;
    vertical-align: top;
  }}
  th {{
    background-color: #f2f4f8;
    font-weight: 600;
    color: #1f2328;
  }}
  tr:nth-child(even) {{
    background-color: #f9fbfd;
  }}
  strong {{
    font-weight: 600;
  }}
</style>
</head>
<body>
{body_content}
</body>
</html>
"""
    return html_template


def build_pdf_document(html_content: str, output_path: Path):
    """Use Chrome or Edge headless mode to print HTML into PDF."""
    HTML_TEMP_PATH.write_text(html_content, encoding="utf-8")

    # Candidate browsers
    candidates = [
        r"C:\Program Files (x86)\Microsoft\Edge\Application\msedge.exe",
        r"C:\Program Files\Microsoft\Edge\Application\msedge.exe",
        r"C:\Program Files\Google\Chrome\Application\chrome.exe",
    ]
    browser_bin = None
    for c in candidates:
        if Path(c).exists():
            browser_bin = c
            break

    if not browser_bin:
        raise RuntimeError("No headless Chrome or Edge browser found on system to render PDF.")

    cmd = [
        browser_bin,
        "--headless",
        "--disable-gpu",
        "--run-all-compositor-stages-before-draw",
        "--no-pdf-header-footer",
        f"--print-to-pdf={output_path}",
        str(HTML_TEMP_PATH.resolve()),
    ]
    res = subprocess.run(cmd, capture_output=True, text=True)
    if HTML_TEMP_PATH.exists():
        HTML_TEMP_PATH.unlink()

    if output_path.exists() and os.path.getsize(output_path) > 0:
        print(f"Generated PDF: {output_path} ({os.path.getsize(output_path):,} bytes)")
    else:
        raise RuntimeError(f"PDF creation failed: {res.stderr or res.stdout}")


def main():
    if not DOC_MD_PATH.exists():
        print(f"Error: {DOC_MD_PATH} not found.")
        sys.exit(1)

    md_text = DOC_MD_PATH.read_text(encoding="utf-8")
    print(f"Reading Markdown source: {DOC_MD_PATH} ({len(md_text):,} chars)")

    # 1. Build Word DOCX
    build_word_document(md_text, DOCX_OUT_PATH)

    # 2. Build PDF
    html_content = build_html_for_pdf(md_text)
    build_pdf_document(html_content, PDF_OUT_PATH)

    print("\n[SUCCESS] Exported documentation successfully in both formats:")
    print(f"  1. DOCX: {DOCX_OUT_PATH}")
    print(f"  2. PDF:  {PDF_OUT_PATH}")


if __name__ == "__main__":
    main()
